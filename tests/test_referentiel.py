"""Regression tests for the corpus referentiel (src/aria_rag/referentiel.py) —
piece-level metadata manifest layered on top of corpus_mapping.yaml.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from aria_rag.config import Settings
from aria_rag.referentiel import (
    Piece,
    PieceFile,
    PieceNotServableError,
    Referentiel,
    export_xlsx,
    import_xlsx,
    load_referentiel,
    match_piece,
    piece_present,
    regenerate_files,
    resolve_document_path,
    save_referentiel,
)


def _piece(**overrides) -> Piece:
    defaults = dict(
        id="p1", official_name="Pièce 1", status_opposabilite="opposable",
        family="annexes", norm_level="local", expected=True, file_match="Annexes/Plans SUP",
        notes="",
    )
    defaults.update(overrides)
    return Piece(**defaults)


# ---------------------------------------------------------------------------
# match_piece — longest-prefix-wins
# ---------------------------------------------------------------------------

def test_match_piece_longest_prefix_wins():
    pieces = [
        _piece(id="plans-sup", file_match="Annexes/Plans SUP"),
        _piece(id="ppri", file_match="Annexes/Plans SUP/PPRI"),
    ]
    assert match_piece("Annexes/Plans SUP/PPRI/ALEA1910.pdf", pieces).id == "ppri"
    assert match_piece("Annexes/Plans SUP/ASUP1.pdf", pieces).id == "plans-sup"


def test_match_piece_no_match_returns_none():
    pieces = [_piece(file_match="Annexes/Plans SUP")]
    assert match_piece("OAP/Sectorielles/foo.pdf", pieces) is None


def test_match_piece_ignores_expected_absent_pieces():
    pieces = [_piece(file_match=None)]
    assert match_piece("anything.pdf", pieces) is None


def test_piece_present_true_when_a_file_maps_to_it():
    files = [PieceFile(path="a.pdf", piece_id="p1", validity="current", chunk_count=3, pages=2)]
    assert piece_present("p1", files) is True
    assert piece_present("p2", files) is False


# ---------------------------------------------------------------------------
# regenerate_files
# ---------------------------------------------------------------------------

def test_regenerate_files_matches_and_marks_validity(tmp_path):
    docs_dir = tmp_path / "Ressources"
    (docs_dir / "Annexes" / "Plans SUP").mkdir(parents=True)
    (docs_dir / "Annexes" / "Plans SUP" / "ASUP1.pdf").write_bytes(b"%PDF-1.4")
    settings = Settings(docs_dir=docs_dir, index_dir=tmp_path / "index")
    settings.index_dir.mkdir(parents=True, exist_ok=True)

    pieces = [_piece(file_match="Annexes/Plans SUP")]
    files = regenerate_files(settings, pieces, mapping_rules=[])
    assert len(files) == 1
    assert files[0].piece_id == "p1"
    assert files[0].validity == "current"
    assert files[0].chunk_count == 0  # no chunks.json yet


# ---------------------------------------------------------------------------
# resolve_document_path
# ---------------------------------------------------------------------------

def test_resolve_document_path_single_current_file(tmp_path):
    docs_dir = tmp_path / "Ressources"
    docs_dir.mkdir()
    referentiel = Referentiel(
        pieces=[_piece(id="reg-ecrit-t1")],
        files=[PieceFile(path="Tome1/REG1_MS1.pdf", piece_id="reg-ecrit-t1", validity="current", chunk_count=10, pages=250)],
    )
    resolved = resolve_document_path("reg-ecrit-t1", docs_dir, referentiel)
    assert resolved == (docs_dir / "Tome1" / "REG1_MS1.pdf").resolve()


def test_resolve_document_path_zero_files_raises(tmp_path):
    docs_dir = tmp_path / "Ressources"
    docs_dir.mkdir()
    referentiel = Referentiel(pieces=[_piece(id="annexe-sanitaire", file_match=None)], files=[])
    try:
        resolve_document_path("annexe-sanitaire", docs_dir, referentiel)
        assert False, "expected PieceNotServableError"
    except PieceNotServableError:
        pass


def test_resolve_document_path_multi_file_raises(tmp_path):
    docs_dir = tmp_path / "Ressources"
    docs_dir.mkdir()
    referentiel = Referentiel(
        pieces=[_piece(id="atlas1")],
        files=[
            PieceFile(path="a.pdf", piece_id="atlas1", validity="current", chunk_count=1, pages=1),
            PieceFile(path="b.pdf", piece_id="atlas1", validity="current", chunk_count=1, pages=1),
        ],
    )
    try:
        resolve_document_path("atlas1", docs_dir, referentiel)
        assert False, "expected PieceNotServableError"
    except PieceNotServableError:
        pass


def test_resolve_document_path_ignores_superseded_files(tmp_path):
    """A superseded twin must not count toward the single-current-file case."""
    docs_dir = tmp_path / "Ressources"
    docs_dir.mkdir()
    referentiel = Referentiel(
        pieces=[_piece(id="reg-ecrit-t1")],
        files=[
            PieceFile(path="REG1.pdf", piece_id="reg-ecrit-t1", validity="superseded", chunk_count=0, pages=250),
            PieceFile(path="REG1_MS1.pdf", piece_id="reg-ecrit-t1", validity="current", chunk_count=10, pages=250),
        ],
    )
    resolved = resolve_document_path("reg-ecrit-t1", docs_dir, referentiel)
    assert resolved.name == "REG1_MS1.pdf"


# ---------------------------------------------------------------------------
# YAML save/load round trip
# ---------------------------------------------------------------------------

def test_save_and_load_referentiel_round_trip(tmp_path):
    path = tmp_path / "referentiel.yaml"
    referentiel = Referentiel(
        pieces=[_piece(id="p1"), _piece(id="p2", file_match=None, expected=False)],
        files=[PieceFile(path="a.pdf", piece_id="p1", validity="current", chunk_count=2, pages=5)],
    )
    save_referentiel(referentiel, path)
    loaded = load_referentiel(path)
    assert loaded.pieces == referentiel.pieces
    assert loaded.files == referentiel.files


# ---------------------------------------------------------------------------
# xlsx export / import round trip
# ---------------------------------------------------------------------------

def test_xlsx_export_import_round_trip(tmp_path):
    docs_dir = tmp_path / "Ressources"
    (docs_dir / "Annexes" / "Plans SUP").mkdir(parents=True)
    (docs_dir / "Annexes" / "Plans SUP" / "ASUP1.pdf").write_bytes(b"%PDF-1.4")
    settings = Settings(docs_dir=docs_dir, index_dir=tmp_path / "index")
    settings.index_dir.mkdir(parents=True, exist_ok=True)

    pieces = [
        _piece(id="plans-sup", file_match="Annexes/Plans SUP"),
        _piece(id="annexe-sanitaire", file_match=None, notes="à confirmer"),
    ]
    files = regenerate_files(settings, pieces, mapping_rules=[])
    original = Referentiel(pieces=pieces, files=files)

    xlsx_path = tmp_path / "referentiel_test.xlsx"
    export_xlsx(original, xlsx_path)

    reimported_pieces = import_xlsx(xlsx_path, original)
    reimported_files = regenerate_files(settings, reimported_pieces, mapping_rules=[])
    roundtripped = Referentiel(pieces=reimported_pieces, files=reimported_files)

    assert roundtripped.pieces == original.pieces
    assert roundtripped.files == original.files


def test_xlsx_import_preserves_file_match_but_updates_status(tmp_path):
    """file_match is structural (engineer-owned) — must survive even though
    it is never shown/edited in the spreadsheet; status_opposabilite is
    the annotable field and must pick up an edit.
    """
    docs_dir = tmp_path / "Ressources"
    docs_dir.mkdir(parents=True)
    settings = Settings(docs_dir=docs_dir, index_dir=tmp_path / "index")
    settings.index_dir.mkdir(parents=True, exist_ok=True)

    original = Referentiel(pieces=[_piece(id="p1", status_opposabilite="opposable")], files=[])
    xlsx_path = tmp_path / "referentiel_test.xlsx"
    export_xlsx(original, xlsx_path)

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    ws = wb["Pièces"]
    ws["C2"] = "informatif"  # edit status_opposabilite
    wb.save(xlsx_path)

    pieces = import_xlsx(xlsx_path, original)
    assert pieces[0].status_opposabilite == "informatif"
    assert pieces[0].file_match == "Annexes/Plans SUP"

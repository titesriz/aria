"""Corpus referentiel — piece-level metadata manifest for the official PLU
dossier structure, layered on top of corpus_mapping.yaml's file-level
family/norm_level/city/validity classification.

Where corpus_mapping.yaml answers "what family is this file", referentiel.yaml
answers "what official dossier piece does this file belong to, and is it
opposable" — the domain-expert-facing (Charline) layer, round-tripped through
Excel (see export_xlsx/import_xlsx) rather than edited as YAML directly.

Two levels, both stored in referentiel.yaml:
  - pieces: hand-authored/seeded. One entry per document an architect would
    recognize as a distinct piece of the PLU dossier (a tome, an atlas, a
    named annexe group, an OAP group, PADD, the rapport de présentation, the
    CCH) — including pieces expected by the Code de l'urbanisme's official
    PLU dossier structure but absent from today's corpus (expected=true,
    file_match=null).
  - files: GENERATED, never hand-edited — see regenerate_files(). Persisted
    to referentiel.yaml anyway (git-tracked) so the file->piece mapping is
    reviewable in diffs and so the /document/{piece_id} endpoint can resolve
    a piece's primary PDF without rescanning the whole corpus on every
    request.
"""
from __future__ import annotations

import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pypdf import PdfReader

from aria_rag.config import ROOT_DIR, Settings
from aria_rag.corpus_mapping import MappingRule, classify_path, load_rules
from aria_rag.indexer import load_existing_chunks
from aria_rag.loader import iter_pdf_paths

DEFAULT_REFERENTIEL_PATH = ROOT_DIR / "referentiel.yaml"

STATUS_OPPOSABILITE_VALUES = (
    "opposable",
    "opposable_compatibilite",
    "informatif",
    "norme_nationale",
)

_HEADER = """\
# Corpus referentiel — piece-level metadata manifest for the official PLU
# dossier structure. Read/written by aria_rag.referentiel. See that module's
# docstring for the schema. Round-trip this file via
# `aria-rag referentiel export` / `aria-rag referentiel import <xlsx>` rather
# than hand-editing the `files:` section (regenerated every import/export —
# any hand edit there is silently discarded).
#
# pieces: hand-authored. file_match is a folder/file PREFIX, relative to
# Settings.docs_dir, forward-slash, NFC-normalized — same longest-prefix-wins
# matching convention as corpus_mapping.yaml. null for a piece expected by
# the official dossier structure but absent from today's corpus.
#
# files: GENERATED — do not hand-edit. Regenerated from the current corpus
# (on-disk PDFs + data/index/chunks.json) every export/import.
"""


@dataclass(slots=True)
class Piece:
    id: str
    official_name: str
    status_opposabilite: str
    family: str | None
    norm_level: str | None
    expected: bool
    file_match: str | None
    notes: str = ""


@dataclass(slots=True)
class PieceFile:
    path: str
    piece_id: str | None
    validity: str
    chunk_count: int
    pages: int | None


@dataclass(slots=True)
class Referentiel:
    pieces: list[Piece] = field(default_factory=list)
    files: list[PieceFile] = field(default_factory=list)


class PieceNotServableError(Exception):
    """A piece exists in the referentiel but has no single current file to
    serve — either zero (an expected-but-absent piece, or every file
    superseded) or more than one (a true multi-file piece, e.g. an atlas).
    """


def load_referentiel(path: Path | None = None) -> Referentiel:
    path = path or DEFAULT_REFERENTIEL_PATH
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    pieces = [Piece(**p) for p in data.get("pieces", [])]
    files = [PieceFile(**f) for f in data.get("files", [])]
    return Referentiel(pieces=pieces, files=files)


def save_referentiel(referentiel: Referentiel, path: Path | None = None) -> None:
    path = path or DEFAULT_REFERENTIEL_PATH
    payload = {
        "pieces": [asdict(p) for p in referentiel.pieces],
        "files": [asdict(f) for f in referentiel.files],
    }
    body = yaml.safe_dump(payload, allow_unicode=True, sort_keys=False, default_flow_style=False)
    path.write_text(_HEADER + "\n" + body, encoding="utf-8")


def match_piece(rel_path: str, pieces: list[Piece]) -> Piece | None:
    """Longest-matching-prefix, same convention as corpus_mapping.classify_path."""
    best: Piece | None = None
    for piece in pieces:
        if piece.file_match is None:
            continue
        if rel_path.startswith(piece.file_match) and (best is None or len(piece.file_match) > len(best.file_match)):
            best = piece
    return best


def piece_present(piece_id: str, files: list[PieceFile]) -> bool:
    return any(f.piece_id == piece_id for f in files)


def _relative_posix(path: Path, docs_dir: Path) -> str:
    try:
        rel = path.resolve().relative_to(docs_dir.resolve()).as_posix()
    except ValueError:
        rel = path.as_posix()
    return unicodedata.normalize("NFC", rel)


def regenerate_files(
    settings: Settings,
    pieces: list[Piece],
    mapping_rules: list[MappingRule] | None = None,
) -> list[PieceFile]:
    """Recomputes the files: table from the current corpus on disk plus the
    current index (data/index/chunks.json). chunk_count is 0 (never None)
    for a superseded or not-yet-indexed file — "zero chunks" is a real,
    meaningful value here, distinct from "not indexed yet", which this
    function has no way to tell apart from the former anyway. pages is None
    only if the PDF could not be opened (matches corpus reality: a handful
    of files elsewhere in this pipeline are known-fragile — see
    checks/known_zero_chunk_files.json).
    """
    rules = mapping_rules if mapping_rules is not None else load_rules()
    chunks_by_source = load_existing_chunks(settings.index_dir)

    entries: list[PieceFile] = []
    for path in iter_pdf_paths(settings.docs_dir):
        rel = _relative_posix(path, settings.docs_dir)
        classification = classify_path(path, settings.docs_dir, rules)
        piece = match_piece(rel, pieces)
        chunk_count = len(chunks_by_source.get(str(path), []))
        try:
            pages = len(PdfReader(str(path)).pages)
        except Exception:
            pages = None
        entries.append(PieceFile(
            path=rel,
            piece_id=piece.id if piece else None,
            validity=classification.validity,
            chunk_count=chunk_count,
            pages=pages,
        ))
    entries.sort(key=lambda e: e.path)
    return entries


def resolve_document_path(piece_id: str, docs_dir: Path, referentiel: Referentiel) -> Path:
    """Resolves a piece's single current PDF for /document/{piece_id}.

    Path safety: the returned path is always built from a `files:` entry
    already computed by regenerate_files() from an on-disk scan — piece_id
    (the only piece of user input in this path) is used solely as a lookup
    key, never concatenated into a filesystem path.
    """
    current_files = [f for f in referentiel.files if f.piece_id == piece_id and f.validity == "current"]
    if not current_files:
        raise PieceNotServableError(f"Aucun fichier courant pour la pièce '{piece_id}'.")
    if len(current_files) > 1:
        raise PieceNotServableError(
            f"La pièce '{piece_id}' comporte {len(current_files)} fichiers "
            "(ex. atlas multi-planches) — non servie individuellement pour le moment."
        )
    docs_dir_resolved = docs_dir.resolve()
    resolved = (docs_dir_resolved / current_files[0].path).resolve()
    if resolved != docs_dir_resolved and docs_dir_resolved not in resolved.parents:
        raise PieceNotServableError("Chemin de fichier hors du corpus — refusé.")
    return resolved


# ---------------------------------------------------------------------------
# Excel export / import — the domain-expert-facing round trip.
# ---------------------------------------------------------------------------

_PIECES_SHEET = "Pièces"
_FILES_SHEET = "Fichiers"

# Column order for the "Pièces" sheet. "id" is a stable join key (hidden-ish,
# first column) — never re-derived from official_name, which is free text.
_PIECES_HEADERS = [
    "ID pièce", "Nom officiel", "Statut d'opposabilité", "Famille",
    "Niveau de norme", "Attendue", "Présente", "Nb fichiers", "Notes",
]
_FILES_HEADERS = ["Chemin", "Pièce", "Validité", "Nb chunks", "Pages"]

_TRUE_STRINGS = {"true", "vrai", "1", "oui", "x"}


def _bool_cell(value: bool) -> str:
    return "VRAI" if value else "FAUX"


def _parse_bool_cell(value) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in _TRUE_STRINGS


def default_export_path(as_of: date | None = None) -> Path:
    as_of = as_of or date.today()
    return ROOT_DIR / f"referentiel_{as_of:%Y%m%d}.xlsx"


def export_xlsx(referentiel: Referentiel, out_path: Path) -> Path:
    """Sheet 1 "Pièces" is the annotable one (status dropdown, wide Notes
    column) — meant for a domain expert (Charline), not an engineer. Sheet 2
    "Fichiers" is generated reference material, included for context but not
    read back by import_xlsx.
    """
    file_counts: dict[str, int] = {}
    for f in referentiel.files:
        if f.piece_id is not None:
            file_counts[f.piece_id] = file_counts.get(f.piece_id, 0) + 1

    wb = Workbook()
    ws1 = wb.active
    ws1.title = _PIECES_SHEET
    ws1.append(_PIECES_HEADERS)
    ws1.freeze_panes = "A2"
    for p in referentiel.pieces:
        ws1.append([
            p.id,
            p.official_name,
            p.status_opposabilite,
            p.family or "",
            p.norm_level or "",
            _bool_cell(p.expected),
            _bool_cell(piece_present(p.id, referentiel.files)),
            file_counts.get(p.id, 0),
            p.notes,
        ])

    status_col = "C"
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(STATUS_OPPOSABILITE_VALUES)}"',
        allow_blank=False,
        showDropDown=False,  # openpyxl quirk: False is what actually shows the dropdown arrow
    )
    ws1.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{ws1.max_row}")

    widths = {"A": 28, "B": 46, "C": 24, "D": 20, "E": 16, "F": 10, "G": 10, "H": 10, "I": 70}
    for col, width in widths.items():
        ws1.column_dimensions[col].width = width

    ws2 = wb.create_sheet(_FILES_SHEET)
    ws2.append(_FILES_HEADERS)
    ws2.freeze_panes = "A2"
    for f in sorted(referentiel.files, key=lambda e: e.path):
        ws2.append([f.path, f.piece_id or "", f.validity, f.chunk_count, f.pages])
    ws2_widths = {"A": 90, "B": 30, "C": 12, "D": 10, "E": 8}
    for col, width in ws2_widths.items():
        ws2.column_dimensions[col].width = width

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def import_xlsx(xlsx_path: Path, existing: Referentiel) -> list[Piece]:
    """Reads only the hand-editable fields (official_name, status, family,
    norm_level, expected, notes) from the "Pièces" sheet; file_match — a
    structural field an engineer sets via corpus_mapping-style prefixes, not
    something a domain expert edits in Excel — is preserved from `existing`
    by id. present/Nb fichiers columns are ignored (generated, recomputed by
    the caller via regenerate_files()).
    """
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[_PIECES_SHEET]
    existing_by_id = {p.id: p for p in existing.pieces}

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in _PIECES_HEADERS}

    pieces: list[Piece] = []
    for row in rows:
        if row[idx["ID pièce"]] in (None, ""):
            continue
        piece_id = str(row[idx["ID pièce"]])
        old = existing_by_id.get(piece_id)
        pieces.append(Piece(
            id=piece_id,
            official_name=str(row[idx["Nom officiel"]] or ""),
            status_opposabilite=str(row[idx["Statut d'opposabilité"]] or ""),
            family=(str(row[idx["Famille"]]) if row[idx["Famille"]] not in (None, "") else None),
            norm_level=(str(row[idx["Niveau de norme"]]) if row[idx["Niveau de norme"]] not in (None, "") else None),
            expected=_parse_bool_cell(row[idx["Attendue"]]),
            file_match=old.file_match if old else None,
            notes=str(row[idx["Notes"]] or ""),
        ))
    return pieces
